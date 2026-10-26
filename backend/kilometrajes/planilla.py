"""Una misma matriz para la pantalla, el Excel y el PDF. Sin alterar catálogos."""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from xml.sax.saxutils import escape

from .models import DestinoKilometraje


def columnas_actuales():
    orden = ("gatica", "martins", "tunuy", "aibak", "marthos", "sharp", "muedaz", "porro", "cruz")
    def posicion(d):
        return next((i for i, nombre in enumerate(orden) if nombre in d.nombre.casefold()), len(orden)), d.nombre
    return [{"id": d.pk, "nombre": d.nombre, "distancia_km": str(d.distancia_km), "activo": True}
            for d in sorted(DestinoKilometraje.objects.filter(activo=True), key=posicion)]


def columnas_control(control):
    if control.columnas:
        return control.columnas
    # Los controles anteriores no tenían instantánea: preservar sus distancias registradas.
    columnas = {d["id"]: d for d in columnas_actuales()}
    for r in control.registros.all():
        columnas[r.destino_id] = {"id": r.destino_id, "nombre": r.destino.nombre,
                                  "distancia_km": str(r.distancia_km), "activo": r.destino.activo}
    return list(columnas.values())


def matriz(control):
    columnas = columnas_control(control)
    registros = {(r.fecha, r.destino_id): r for r in control.registros.all()}
    filas, totales = [], [Decimal("0")] * len(columnas)
    fecha = control.fecha_desde
    while fecha <= control.fecha_hasta:
        valores = []
        for i, d in enumerate(columnas):
            registro = registros.get((fecha, d["id"]))
            valores.append(registro.cantidad_viajes if registro else None)
            if registro:
                totales[i] += registro.kilometros
        filas.append((fecha, valores))
        fecha += timedelta(days=1)
    return columnas, filas, totales


def titulo(control):
    return f"CONTROL DE KILOMETRAJE REPARTO | Período: {control.fecha_desde:%d/%m/%Y} al {control.fecha_hasta:%d/%m/%Y}"


def excel(control):
    columnas, filas, totales = matriz(control)
    w = Workbook()
    s = w.active
    s.title = "Control kilometraje"
    n = len(columnas) + 1
    s.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    s.cell(1, 1, f"CHOFER\n{control.chofer.nombre}")
    s.merge_cells(start_row=1, start_column=2, end_row=2, end_column=n)
    s.cell(1, 2, titulo(control))
    s.cell(3, 1, "DESTINO")
    for col, destino in enumerate(columnas, 2):
        s.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        s.cell(3, col, destino["nombre"])
        s.cell(5, col, float(destino["distancia_km"])).number_format = '0.##" km"'
        s.column_dimensions[get_column_letter(col)].width = 17
    s.cell(5, 1, "DISTANCIA")
    s.cell(6, 1, "DÍA")
    gris = Side(style="thin", color="DAE2EF")
    for indice, (fecha, valores) in enumerate(filas, 7):
        s.cell(indice, 1, fecha).number_format = "d"
        for col, valor in enumerate(valores, 2):
            s.cell(indice, col, valor)
        if fecha.weekday() == 6:
            for celda in s[indice][:n]:
                celda.fill = PatternFill("solid", fgColor="858585")
                celda.font = Font(color="FFFFFF")
    total_fila = 7 + len(filas)
    s.cell(total_fila, 1, "TOTAL KMS")
    registros = {(r.fecha, r.destino_id): r for r in control.registros.all()}
    for col, total in enumerate(totales, 2):
        letra = get_column_letter(col)
        destino = columnas[col-2]
        distancias = [registros[(fecha, destino["id"])].distancia_km if (fecha, destino["id"]) in registros else Decimal(destino["distancia_km"]) for fecha, _ in filas]
        if any(d != Decimal(destino["distancia_km"]) for d in distancias):
            factores = ";".join(str(d) for d in distancias)
            formula = f"=SUMPRODUCT({letra}7:{letra}{total_fila-1},{{{factores}}})"
        else:
            formula = f"=SUM({letra}7:{letra}{total_fila-1})*{letra}5"
        s.cell(total_fila, col, formula).number_format = '0.##" km"'
    final = total_fila + 2
    s.merge_cells(start_row=final, start_column=1, end_row=final, end_column=max(1, n-2))
    s.cell(final, 1, "TOTAL KILÓMETROS RECORRIDOS EN EL PERÍODO:")
    s.merge_cells(start_row=final, start_column=n-1, end_row=final, end_column=n)
    s.cell(final, n-1, f"=SUM(B{total_fila}:{get_column_letter(n)}{total_fila})").number_format = '0.##" km"'
    for row in s.iter_rows(min_row=1, max_row=final, max_col=n):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(left=gris, right=gris, top=gris, bottom=gris)
            if c.row <= 4 or c.row == 6:
                c.fill = PatternFill("solid", fgColor="243247")
                c.font = Font(color="FFFFFF", bold=True, size=10)
            elif c.row == 5:
                c.fill = PatternFill("solid", fgColor="E4EAF3")
                c.font = Font(bold=True)
            elif c.row == total_fila:
                c.font = Font(color="E65A00", bold=True)
            elif c.row == final:
                c.fill = PatternFill("solid", fgColor="107C78")
                c.font = Font(color="FFFFFF", bold=True)
    s.column_dimensions["A"].width = 18
    s.row_dimensions[1].height = 25
    s.row_dimensions[3].height = 28
    for r in range(7, total_fila):
        s.row_dimensions[r].height = 20
    s.row_dimensions[final].height = 30
    s.freeze_panes = "B7"
    s.print_title_rows = "1:6"
    s.print_area = f"A1:{get_column_letter(n)}{final}"
    s.page_setup.orientation = "landscape"
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
    validacion = DataValidation(type="whole", operator="between", formula1=0, formula2=2147483647, allow_blank=True)
    validacion.errorTitle = "Cantidad inválida"
    validacion.error = "Ingresá una cantidad entera de viajes, mayor o igual a cero."
    validacion.showErrorMessage = True
    s.add_data_validation(validacion)
    validacion.add(f"B7:{get_column_letter(n)}{total_fila-1}")
    if control.observaciones:
        s.cell(final+2, 1, "Observaciones")
        # Forzar texto: los datos del usuario nunca se interpretan como fórmulas.
        c = s.cell(final+2, 2, control.observaciones)
        c.data_type = "s"
        s.print_area = f"A1:{get_column_letter(n)}{final+2}"
    for dcol in range(2, n+1):
        s.cell(3, dcol).data_type = "s"
    s.cell(1, 1).data_type = "s"
    salida = BytesIO()
    w.save(salida)
    return salida.getvalue()


def pdf(control):
    columnas, filas, totales = matriz(control)
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=18, bottomMargin=18)
    estilo = getSampleStyleSheet()["Normal"]
    estilo.fontSize = 8
    estilo.leading = 10
    cabecera_estilo = ParagraphStyle("cabecera", parent=estilo, textColor=colors.white, fontSize=8, leading=9, alignment=1)
    ancho = (landscape(A4)[0] - 80) / max(1, len(columnas))
    cabecera = ["DÍA"] + [Paragraph(escape(d["nombre"]), cabecera_estilo) for d in columnas]
    datos = [cabecera, ["DISTANCIA"] + [f'{d["distancia_km"]} km' for d in columnas]]
    datos += [[fecha.strftime("%d/%m")] + [v if v is not None else "" for v in valores] for fecha, valores in filas]
    datos += [["TOTAL KMS"] + [f"{v:g} km" for v in totales]]
    tabla = Table(datos, colWidths=[60] + [ancho]*len(columnas), repeatRows=2, rowHeights=None)
    comandos = [("GRID", (0,0), (-1,-1), .3, colors.HexColor("#d8e0eb")),
                ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#243247")),
                ("TEXTCOLOR", (0,0), (-1,0), colors.white),
                ("BACKGROUND", (0,1), (-1,1), colors.HexColor("#e4eaf3")),
                ("ALIGN", (0,0), (-1,-1), "CENTER"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("FONTSIZE", (0,0), (-1,-1), 8), ("LEADING", (0,0), (-1,-1), 9), ("TOPPADDING", (0,0), (-1,-1), 1.5),
                ("BOTTOMPADDING", (0,0), (-1,-1), 1.5), ("TEXTCOLOR", (0,-1), (-1,-1), colors.HexColor("#d45a00"))]
    for i, (fecha, _) in enumerate(filas, 2):
        if fecha.weekday() == 6:
            comandos += [("BACKGROUND", (0,i), (-1,i), colors.HexColor("#858585")), ("TEXTCOLOR", (0,i), (-1,i), colors.white)]
    tabla.setStyle(TableStyle(comandos))
    total = Table([["TOTAL KILÓMETROS RECORRIDOS EN EL PERÍODO:", f"{sum(totales):g} km"]], colWidths=[doc.width*.8, doc.width*.2])
    total.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor("#107c78")),
        ("TEXTCOLOR", (0,0), (-1,-1), colors.white), ("FONTSIZE", (0,0), (-1,-1), 9),
        ("TOPPADDING", (0,0), (-1,-1), 10), ("BOTTOMPADDING", (0,0), (-1,-1), 10),
        ("ALIGN", (-1,0), (-1,0), "CENTER")]))
    contenido = [Paragraph(escape(titulo(control)), estilo), Spacer(1, 6),
                 Paragraph("CHOFER: " + escape(control.chofer.nombre), estilo), Spacer(1, 6), tabla, Spacer(1, 8),
                 total]
    if control.observaciones:
        contenido += [Spacer(1, 6), Paragraph("Observaciones: " + escape(control.observaciones), estilo)]
    doc.build(contenido)
    return salida.getvalue()
