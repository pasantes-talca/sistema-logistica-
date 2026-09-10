import re
from datetime import datetime

from openpyxl import load_workbook


# ============================================================
# UTILIDADES
# ============================================================

def normalizar_texto(valor):
    return str(
        valor or ""
    ).strip().lower()


def normalizar_carga(valor):

    if valor in (None, ""):
        return ""

    if isinstance(valor, float):

        if valor.is_integer():
            return str(int(valor))

    if isinstance(valor, int):
        return str(valor)

    texto = str(
        valor
    ).strip()

    try:

        numero = float(
            texto
        )

        if numero.is_integer():
            return str(
                int(numero)
            )

    except ValueError:
        pass

    return texto


def es_numero_carga(valor):
    """
    Determina si una celda parece contener
    un N° de carga real.
    """

    texto = normalizar_carga(
        valor
    )

    if not texto:
        return False

    return texto.isdigit()


# ============================================================
# FECHA
# ============================================================

def extraer_fecha(hoja):

    patron = re.compile(
        r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})"
    )

    for fila in hoja.iter_rows(
        min_row=1,
        max_row=min(
            hoja.max_row,
            5,
        ),
        values_only=True,
    ):

        for valor in fila:

            if isinstance(
                valor,
                datetime,
            ):
                return valor.date()

            texto = str(
                valor or ""
            )

            coincidencia = patron.search(
                texto
            )

            if coincidencia:

                dia = int(
                    coincidencia.group(1)
                )

                mes = int(
                    coincidencia.group(2)
                )

                anio = int(
                    coincidencia.group(3)
                )

                return datetime(
                    anio,
                    mes,
                    dia,
                ).date()

    return None


# ============================================================
# ENCABEZADOS
# ============================================================

def buscar_fila_encabezados(
    hoja,
):
    """
    Busca la fila que comienza con N° Carga.
    """

    for numero_fila in range(
        1,
        min(
            hoja.max_row,
            20,
        ) + 1,
    ):

        valor = normalizar_texto(
            hoja.cell(
                numero_fila,
                1,
            ).value
        )

        if (
            "carga" in valor
            and (
                "n°" in valor
                or "nº" in valor
                or "n" in valor
            )
        ):
            return numero_fila

    return None


# ============================================================
# LECTOR PRINCIPAL
# ============================================================

def leer_planilla_salidas(
    archivo,
):

    libro = load_workbook(
        archivo,
        data_only=True,
    )


    hoja_objetivo = None


    for nombre in libro.sheetnames:

        if (
            "SALIDA" in nombre.upper()
            and "CHOFER" in nombre.upper()
        ):

            hoja_objetivo = libro[
                nombre
            ]

            break


    if hoja_objetivo is None:

        hoja = libro.active

    else:

        hoja = hoja_objetivo


    fecha = extraer_fecha(
        hoja
    )


    fila_encabezados = (
        buscar_fila_encabezados(
            hoja
        )
    )


    if fila_encabezados is None:

        raise ValueError(
            (
                "No se encontró la tabla "
                "de salida de choferes."
            )
        )


    registros = []


    for numero_fila in range(
        fila_encabezados + 1,
        hoja.max_row + 1,
    ):

        valor_carga = hoja.cell(
            numero_fila,
            1,
        ).value


        # Cuando aparecen textos como
        # "ESTADO DE UNIDADES..." dejamos
        # de considerarlos repartos.
        if not es_numero_carga(
            valor_carga
        ):
            continue


        carga = normalizar_carga(
            valor_carga
        )


        patente = str(
            hoja.cell(
                numero_fila,
                2,
            ).value
            or ""
        ).strip()


        combustible = str(
            hoja.cell(
                numero_fila,
                3,
            ).value
            or ""
        ).strip()


        chofer = str(
            hoja.cell(
                numero_fila,
                4,
            ).value
            or ""
        ).strip()


        ayudante_1 = str(
            hoja.cell(
                numero_fila,
                5,
            ).value
            or ""
        ).strip()


        ayudante_2 = str(
            hoja.cell(
                numero_fila,
                6,
            ).value
            or ""
        ).strip()


        ayudante_3 = str(
            hoja.cell(
                numero_fila,
                7,
            ).value
            or ""
        ).strip()


        observaciones = str(
            hoja.cell(
                numero_fila,
                8,
            ).value
            or ""
        ).strip()


        ayudantes = [
            nombre
            for nombre in (
                ayudante_1,
                ayudante_2,
                ayudante_3,
            )
            if nombre
        ]


        registros.append(
            {
                "carga":
                    carga,

                "patente":
                    patente,

                "combustible":
                    combustible,

                "chofer":
                    chofer,

                "ayudantes":
                    ayudantes,

                "observaciones":
                    observaciones,
            }
        )


    return {
        "fecha":
            fecha,

        "registros":
            registros,
    }